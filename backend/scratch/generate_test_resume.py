import docx
from pathlib import Path

def create_resume():
    doc = docx.Document()
    
    # Title
    doc.add_heading('John Doe - Senior Software Engineer', 0)
    
    # Profile
    doc.add_heading('Professional Summary', level=1)
    doc.add_paragraph(
        'Experienced Senior Software Engineer with over 5 years of expertise in full-stack web application development. '
        'Proven track record of designing, building, and deploying highly scalable API backends and interactive frontends. '
        'Passionate about clean code, automated testing, containerized workflows, and agile software development practices.'
    )
    
    # Skills
    doc.add_heading('Technical Skills', level=1)
    p = doc.add_paragraph()
    p.add_run('Languages: ').bold = True
    p.add_run('Python, JavaScript, TypeScript, SQL, HTML, CSS\n')
    p.add_run('Frameworks & Tools: ').bold = True
    p.add_run('FastAPI, Django, React, Node.js, Git, Docker, Kubernetes\n')
    p.add_run('Databases & Cloud: ').bold = True
    p.add_run('PostgreSQL, MySQL, Redis, AWS, GCP, CI/CD pipelines')
    
    # Experience
    doc.add_heading('Professional Experience', level=1)
    
    doc.add_heading('Senior Software Engineer - Tech Solutions Inc.', level=2)
    doc.add_paragraph('June 2023 - Present')
    doc.add_paragraph(
        '• Designed and implemented microservices using Python and FastAPI, improving API response times by 35%.\n'
        '• Developed responsive, modular web interfaces using React and TypeScript, boosting client satisfaction.\n'
        '• Configured CI/CD pipelines using GitHub Actions to automate testing and deployments on AWS and Kubernetes.\n'
        '• Containerized the entire dev environment using Docker, streamlining onboarding workflows.'
    )
    
    doc.add_heading('Software Developer - Web Innovations', level=2)
    doc.add_paragraph('Jan 2021 - May 2023')
    doc.add_paragraph(
        '• Developed robust RESTful APIs using Python, Django, and PostgreSQL.\n'
        '• Handled database optimizations, reducing slow query times using indexing and optimized SQL joins.\n'
        '• Collaborated with product teams in an Agile Scrum environment to deliver monthly features.'
    )
    
    # Education
    doc.add_heading('Education', level=1)
    doc.add_paragraph('B.S. in Computer Science - State University (2020)')
    
    # Save the file
    output_path = Path('/app/uploads/test_resume.docx')
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    print(f"Successfully generated test resume at: {output_path}")

if __name__ == '__main__':
    create_resume()
