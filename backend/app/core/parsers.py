"""
Core Resume Parsers.
Handles text extraction from PDF and DOCX files.
"""
import re
from abc import ABC, abstractmethod
from pathlib import Path
import fitz  # PyMuPDF
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
import logging

logger = logging.getLogger(__name__)

class ParserError(Exception):
    """Base exception for parsing errors."""
    pass


class ResumeParser(ABC):
    """Abstract interface for resume parsing."""
    
    @abstractmethod
    def extract_text(self, file_path: Path) -> str:
        """
        Extract raw text from a file.
        
        Args:
            file_path: Path to the document.
            
        Returns:
            Clean extracted text.
            
        Raises:
            FileNotFoundError: If the file does not exist.
            ParserError: If parsing fails (corrupted, encrypted, etc.).
        """
        pass
        
    def _clean_text(self, text: str) -> str:
        """
        Cleans and normalizes extracted raw text:
        - Removes null characters (Postgres safety).
        - Normalizes line breaks to single \n.
        - Limits consecutive newlines to maximum of 2.
        - Strips leading/trailing whitespaces.
        """
        if not text:
            return ""
            
        # Replace null bytes to prevent database write errors
        text = text.replace("\x00", "")
        
        # Normalize Windows/Mac line endings
        text = text.replace("\r\n", "\n").replace("\r", "\n")
        
        # Collapse 3+ consecutive newlines into 2 (blank line preservation)
        text = re.sub(r"\n{3,}", "\n\n", text)
        
        return text.strip()


class PDFResumeParser(ResumeParser):
    """Extracts raw text from a PDF file using PyMuPDF."""
    
    def extract_text(self, file_path: Path) -> str:
        if not file_path.exists():
            raise FileNotFoundError(f"PDF file not found at {file_path}")
            
        text_parts = []
        try:
            # PyMuPDF will raise an exception on completely corrupted files
            doc = fitz.open(str(file_path))
        except Exception as e:
            logger.error(f"Failed to open PDF file {file_path}: {e}")
            raise ParserError(f"Failed to open or parse PDF file. It might be corrupted. Error: {e}")
            
        if doc.is_encrypted:
            # Attempt to authenticate with empty password (some encrypted PDFs allow this for viewing)
            if not doc.authenticate(""):
                doc.close()
                raise ParserError("PDF is encrypted and requires a password to extract text.")
                
        try:
            for page in doc:
                text = page.get_text()
                if text:
                    text_parts.append(text)
        except Exception as e:
            logger.error(f"Error reading pages from PDF {file_path}: {e}")
            raise ParserError(f"Error extracting text from PDF pages. Error: {e}")
        finally:
            doc.close()
            
        raw_text = "\n".join(text_parts)
        cleaned_text = self._clean_text(raw_text)
        
        if not cleaned_text:
            logger.warning(f"Extracted empty text from PDF {file_path}")
            
        return cleaned_text


class DOCXResumeParser(ResumeParser):
    """Extracts raw text from a DOCX file using python-docx."""
    
    def extract_text(self, file_path: Path) -> str:
        if not file_path.exists():
            raise FileNotFoundError(f"DOCX file not found at {file_path}")
            
        try:
            doc = Document(file_path)
        except PackageNotFoundError:
            raise ParserError(f"Failed to open DOCX file. It might be corrupted or not a valid DOCX. {file_path}")
        except Exception as e:
            logger.error(f"Failed to open DOCX file {file_path}: {e}")
            raise ParserError(f"Failed to open DOCX file. Error: {e}")
            
        text_parts = []
        
        try:
            # 1. Paragraphs (body text)
            for paragraph in doc.paragraphs:
                if paragraph.text:
                    text_parts.append(paragraph.text)
                    
            # 2. Tables (structural text)
            for table in doc.tables:
                for row in table.rows:
                    row_cells = []
                    for cell in row.cells:
                        txt = cell.text.strip()
                        if txt and (not row_cells or row_cells[-1] != txt):
                            row_cells.append(txt)
                    if row_cells:
                        text_parts.append(" | ".join(row_cells))
        except Exception as e:
            logger.error(f"Error extracting text from DOCX {file_path}: {e}")
            raise ParserError(f"Error extracting text from DOCX. Error: {e}")
            
        raw_text = "\n".join(text_parts)
        cleaned_text = self._clean_text(raw_text)
        
        if not cleaned_text:
            logger.warning(f"Extracted empty text from DOCX {file_path}")
            
        return cleaned_text
